import hashlib
import io
import json
import re
import sqlite3
import zipfile
from concurrent.futures import ThreadPoolExecutor

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from contractdb.backup import create_backup, restore_backup
from contractdb.folder_export import export_dev_fee_folders
from contractdb.db import connect
from contractdb.extract import extract_pages, suggest_fields
from contractdb.web import create_app


def text_pdf(label="TEST-ONLY"):
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 14 Tf 40 730 Td ({label} - This is a generated acceptance fixture. No real contract or customer data is contained here. This document is used only for local testing.) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.fixture
def app(tmp_path):
    app = create_app({"TESTING": True, "DATA_DIR": tmp_path / "data", "SYNC_EXTRACTION": True, "RESUME_EXTRACTION": False})
    return app


def csrf(client):
    with client.session_transaction() as session:
        return session["csrf"]


def post(client, url, data=None, **kwargs):
    return client.post(url, data={"csrf_token": csrf(client), **(data or {})}, **kwargs)


@pytest.fixture
def client(app):
    client = app.test_client()
    client.get("/setup")
    assert post(client, "/setup", {"username": "testadmin", "password": "acceptance-test-password"}).status_code == 302
    assert post(client, "/login", {"username": "testadmin", "password": "acceptance-test-password"}).status_code == 302
    return client


def upload(client, content=None, name="验收测试.pdf"):
    response = post(client, "/upload", {"files": (io.BytesIO(content or text_pdf()), name)}, content_type="multipart/form-data")
    assert response.status_code == 302
    return response.location.split("/")[-1]


def save(client, identifier, **extra):
    return post(client, f"/documents/{identifier}/save", {"revision": "1", "company": "验收测试专用公司", "type_id": "1", "subtype_id": "1", "project": "验收项目", "title": "验收合同", "contract_number": "TEST-001", "signed_date": "2026-09-23", "verified": "1", **extra})


def test_freeform_classification_and_amount_display(app, client):
    identifier = upload(client)
    assert save(client, identifier, type_name="软件服务", subtype_name="维护与许可", amount="123456.78").status_code == 302
    with connect(app.config['DATA_DIR']) as conn:
        row = conn.execute('SELECT d.type_id, d.subtype_id FROM documents d WHERE id = ?', (identifier,)).fetchone()
        assert conn.execute('SELECT name FROM import_types WHERE id = ?', (row['type_id'],)).fetchone()[0] == '软件服务'
        assert conn.execute('SELECT name FROM subtypes WHERE id = ? AND type_id = ?', (row['subtype_id'], row['type_id'])).fetchone()[0] == '维护与许可'
    detail = client.get(f'/documents/{identifier}').get_data(as_text=True)
    assert 'name="type_name"' in detail and 'name="subtype_name"' in detail
    assert '123,456.78' in detail
    assert '123,456.78' in client.get('/').get_data(as_text=True)
    assert save(client, identifier, revision='2', type_name='', subtype_name='维护').status_code == 400
    assert save(client, identifier, revision='2', type_name='软件服务', subtype_name='维护与许可').status_code == 302
    with connect(app.config['DATA_DIR']) as conn:
        assert conn.execute("SELECT count(*) FROM import_types WHERE name = '软件服务'").fetchone()[0] == 1


def test_ocr_reaches_later_scanned_page_and_amount_stays_a_hint(tmp_path, monkeypatch):
    writer = PdfWriter()
    for _ in range(6):
        writer.add_page(PdfReader(io.BytesIO(text_pdf('TEXT-FIRST'))).pages[0])
    scanned = io.BytesIO()
    Image.new('RGB', (200, 120), 'white').save(scanned, format='PDF')
    writer.add_page(PdfReader(io.BytesIO(scanned.getvalue())).pages[0])
    path = tmp_path / 'mixed.pdf'
    with path.open('wb') as output:
        writer.write(output)
    monkeypatch.setattr('contractdb.extract._ocr_image', lambda image: '合同金额：人民币40万元')
    pages, method = extract_pages(path)
    assert ('第 7 页（OCR）', '合同金额：人民币40万元') in pages
    candidates = suggest_fields(pages, method)
    assert candidates['amount_hint'][0]['value'] == '400000.00'
    assert candidates['amount_hint'][0]['page'] == '第 7 页（OCR）'
    assert candidates['amount_hint'][0]['method'] == 'ocr'
    assert suggest_fields([('第 1 页', '签订日期：2026年2月30日')], 'text')['signed_date'] == []


def test_company_filter_keeps_other_company_out_of_totals_and_export(app, client):
    first = upload(client, text_pdf('COMPANY-FIRST'), '第一份.pdf')
    second = upload(client, text_pdf('COMPANY-SECOND'), '第二份.pdf')
    assert save(client, first, company='甲方测试公司', amount='100').status_code == 302
    assert save(client, second, company='乙方测试公司', amount='900').status_code == 302
    with connect(app.config['DATA_DIR']) as conn:
        company_id = conn.execute("SELECT id FROM companies WHERE name = '甲方测试公司'").fetchone()[0]
    page = client.get(f'/?company={company_id}').get_data(as_text=True)
    assert f'/documents/{first}' in page and f'/documents/{second}' not in page
    assert '100.00' in page and '900.00' not in page
    accounts = client.get(f'/companies?company={company_id}').get_data(as_text=True)
    assert f'/documents/{first}' in accounts and f'/documents/{second}' not in accounts
    workbook = load_workbook(io.BytesIO(client.get(f'/export.xlsx?company={company_id}').data))
    assert workbook['合同清单'].max_row == 2
    assert workbook['合同清单']['A2'].value == '甲方测试公司'
    account_book = load_workbook(io.BytesIO(client.get(f'/companies/export.xlsx?company={company_id}').data))
    assert account_book['公司账款汇总'].max_row == 2
    assert account_book['公司账款汇总']['A2'].value == '甲方测试公司'


def test_dev_fee_folder_export_uses_confirmed_year_and_company(app, client, tmp_path):
    original = text_pdf('FOLDER-CONFIRMED')
    confirmed = upload(client, original, '开发费原件.pdf')
    pending = upload(client, text_pdf('FOLDER-PENDING'), '开发费待核对.pdf')
    assert save(client, confirmed, company='归档测试公司', type_name='合同', subtype_name='开发费',
                signed_date='2025-07-09').status_code == 302
    output = tmp_path / '开发费合同归档'
    result = export_dev_fee_folders(app.config['DATA_DIR'], output)
    assert result == {'exported': 1, 'pending': 1, 'pending_clues': 1}
    copies = list((output / '2025' / '归档测试公司').glob('*.pdf'))
    assert len(copies) == 1 and copies[0].read_bytes() == original
    assert pending not in str(copies[0])
    assert '开发费待核对.pdf' in (output / '待核对合同清单.csv').read_text(encoding='utf-8-sig')
    assert export_dev_fee_folders(app.config['DATA_DIR'], output) == result


def test_admin_restart_requires_csrf_and_managed_server(app, client, monkeypatch):
    calls = []
    monkeypatch.setattr('contractdb.web.subprocess.Popen', lambda *args, **kwargs: calls.append((args, kwargs)))
    assert client.post('/admin/restart').status_code == 400
    assert post(client, '/admin/restart').status_code == 503
    app.config.update(SERVICE_PORT=8012, SERVICE_HOST='127.0.0.1')
    assert post(client, '/admin/restart').status_code == 202
    assert calls[0][0][0][-4:] == ['-Port', '8012', '-ListenAddress', '127.0.0.1']
    with connect(app.config['DATA_DIR']) as conn:
        conn.execute("UPDATE users SET role = 'user'")
        conn.commit()
    assert post(client, '/admin/restart').status_code == 403
    assert len(calls) == 1


def test_dashboard_dual_direction_currency_filters_and_void(app, client):
    identifier = upload(client)
    assert save(client, identifier, amount='1000', payable='600').status_code == 302
    base = {'occurred_on': '2026-09-23', 'description': 'Dashboard test'}
    for kind, amount in [('receipt', '300'), ('payment', '200'), ('invoice', '100')]:
        assert post(client, f'/documents/{identifier}/finance', {**base, 'kind': kind, 'amount': amount, 'reference': kind, 'entry_token': 'dashboard-token-' + kind}).status_code == 302
    from contractdb.db import document_query
    from contractdb.finance import amount_summary
    with connect(app.config['DATA_DIR']) as conn:
        sql, params = document_query({})
        group = amount_summary(conn.execute(sql, params))[0]
        assert (group['receivable'], group['received'], group['unreceived']) == (100000, 30000, 70000)
        assert (group['payable'], group['paid'], group['unpaid']) == (60000, 20000, 40000)
        record = conn.execute("SELECT id FROM finance_records WHERE kind='payment'").fetchone()[0]
    assert '700.00' in client.get('/').get_data(as_text=True)
    assert '400.00' in client.get('/?finance=payable&sort=unpaid').get_data(as_text=True)
    assert '验收合同' not in client.get('/?currency=USD').get_data(as_text=True)
    assert post(client, f'/finance/{record}/void', {'reason': 'Dashboard test correction'}).status_code == 302
    assert '600.00' in client.get('/?finance=payable').get_data(as_text=True)
    assert client.get('/?from=2026-10-01&to=2026-09-01').status_code == 400
    assert client.get('/?currency=INVALID').status_code == 400


def test_dashboard_totals_no_netting_or_page_limit():
    from contractdb.finance import amount_summary
    rows = [dict(currency='CNY', amount_minor=10000, payable_minor=5000, received_minor=0, paid_minor=0, invoiced_minor=0) for _ in range(26)]
    rows += [dict(currency='CNY', amount_minor=10000, payable_minor=None, received_minor=20000, paid_minor=0, invoiced_minor=11000),
             dict(currency='USD', amount_minor=None, payable_minor=300, received_minor=0, paid_minor=400, invoiced_minor=0)]
    groups = {g['currency']: g for g in amount_summary(rows)}
    assert groups['CNY']['unreceived'] == 260000
    assert groups['CNY']['excess_receipt'] == 10000
    assert groups['CNY']['excess_invoice'] == 1000
    assert groups['CNY']['unknown_payable'] == 1
    assert groups['USD']['unknown_receivable'] == 1
    assert groups['USD']['excess_payment'] == 100


def test_setup_login_permissions_and_csrf(app, client):
    assert client.get("/").status_code == 200
    assert client.post("/logout").status_code == 400
    post(client, "/admin/users", {"username": "testuser", "password": "ordinary-user-password", "role": "user"})
    post(client, "/logout")
    assert client.get("/").status_code == 302
    client.get("/login")
    post(client, "/login", {"username": "testuser", "password": "ordinary-user-password"})
    assert client.get("/admin").status_code == 403
    assert post(client, "/admin/taxonomy", {"kind": "type", "name": "测试"}).status_code == 403
    assert client.get("/upload").status_code == 200


def test_upload_verify_query_export_download(app, client):
    original = text_pdf()
    identifier = upload(client, original)
    page = client.get(f"/documents/{identifier}")
    assert page.status_code == 200
    assert "待核对" in page.get_data(as_text=True)
    assert save(client, identifier).status_code == 302
    listing = client.get("/?q=TEST-001&company=1&type=1&subtype=1&project=1&status=verified")
    assert "验收合同" in listing.get_data(as_text=True)
    assert client.get(f"/documents/{identifier}/file").data == original
    response = client.get("/export.xlsx?q=TEST-001&status=verified")
    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.active.max_row == 2
    assert workbook.active["A2"].value == "验收测试专用公司"
    assert workbook.active["H2"].value == "已核对"
    response = client.get("/export.zip?q=TEST-001")
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        original_name = next(name for name in archive.namelist() if name.endswith(".pdf"))
        assert archive.read(original_name) == original
        assert "合同清单.xlsx" in archive.namelist()
    empty = load_workbook(io.BytesIO(client.get("/export.xlsx?q=NOT-FOUND").data))
    assert empty.active.max_row == 1


def test_duplicate_invalid_files_and_unknown_fields(app, client):
    identifier = upload(client)
    assert upload(client, name="另一个名称.pdf") == identifier
    bad = post(client, "/upload", {"files": (io.BytesIO(b"not a pdf"), "坏文件.pdf")}, content_type="multipart/form-data")
    assert bad.status_code == 200
    with connect(app.config["DATA_DIR"]) as conn:
        rows = conn.execute("SELECT * FROM documents").fetchall()
        assert len(rows) == 1
        assert rows[0]["company_id"] is None and rows[0]["project_id"] is None and rows[0]["subtype_id"] is None
        assert rows[0]["review_status"] == "pending"


def test_concurrent_edit_archive_and_restore(app, client):
    identifier = upload(client)
    assert save(client, identifier).status_code == 302
    assert save(client, identifier, title="过期修改").status_code == 409
    assert post(client, f"/documents/{identifier}/archive").status_code == 302
    assert "验收合同" not in client.get("/").get_data(as_text=True)
    assert "验收合同" in client.get("/?archived=1").get_data(as_text=True)
    assert post(client, f"/documents/{identifier}/restore").status_code == 302
    assert client.get(f"/documents/{identifier}/file").status_code == 200


def test_word_candidates_require_manual_acceptance(app, client):
    doc = Document()
    doc.add_paragraph("验收测试专用文档，不用于真实业务")
    doc.add_paragraph("甲方：验收测试有限公司")
    doc.add_paragraph("合同编号：WORD-TEST-001")
    doc.add_paragraph("项目名称：测试项目")
    buffer = io.BytesIO()
    doc.save(buffer)
    identifier = upload(client, buffer.getvalue(), "测试文档.docx")
    with connect(app.config["DATA_DIR"]) as conn:
        row = conn.execute("SELECT * FROM documents WHERE id = ?", (identifier,)).fetchone()
        candidates = json.loads(row["candidates_json"])
        assert candidates["contract_number"][0]["value"] == "WORD-TEST-001"
        assert candidates["company"][0]["page"] == "Word 正文"
        assert row["company_id"] is None
        assert row["review_status"] == "pending"


def test_excel_values_cannot_become_formulas(client):
    identifier = upload(client)
    save(client, identifier, title="=SUM(1,2)")
    sheet = load_workbook(io.BytesIO(client.get("/export.xlsx").data)).active
    assert sheet["E2"].value == "=SUM(1,2)"
    assert sheet["E2"].data_type == "s"


def test_last_admin_and_taxonomy_constraints(client):
    assert post(client, "/admin/users/1", {"role": "user", "active": "1"}).status_code == 400
    post(client, "/admin/taxonomy", {"kind": "type", "name": "验收资料"})
    post(client, "/admin/taxonomy", {"kind": "subtype", "name": "验收单", "type_id": "2"})
    identifier = upload(client)
    assert save(client, identifier, type_id="1", subtype_id="2").status_code == 400


def test_backup_restore_and_hash_preservation(app, client, tmp_path):
    original = text_pdf()
    identifier = upload(client, original)
    save(client, identifier)
    backup = create_backup(app.config["DATA_DIR"])
    target = tmp_path / "restored"
    restore_backup(backup, target)
    with connect(target) as conn:
        row = conn.execute("SELECT * FROM documents WHERE id = ?", (identifier,)).fetchone()
        assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    assert (target / "originals" / row["stored_filename"]).read_bytes() == original
    with pytest.raises(ValueError, match="恢复目录"):
        restore_backup(backup, target)


def test_multiuser_parallel_reads_and_uploads(app, client):
    def worker(index):
        with app.test_client() as local:
            local.get("/login")
            with local.session_transaction() as session:
                session["user_id"] = 1
                session["session_version"] = 1
            identifier = upload(local, text_pdf(f"PARALLEL-{index}"), f"并发验收{index}.pdf")
            return local.get(f"/documents/{identifier}").status_code
    with ThreadPoolExecutor(max_workers=4) as executor:
        assert list(executor.map(worker, range(6))) == [200] * 6
    with connect(app.config["DATA_DIR"]) as conn:
        assert conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 6


def test_money_partial_receipts_invoices_void_and_company_report(app, client):
    identifier = upload(client)
    assert save(client, identifier, amount="1000.01", currency="CNY").status_code == 302
    receipt = {"kind": "receipt", "amount": "300.01", "occurred_on": "2026-09-23", "reference": "TEST-RECEIPT", "description": "验收测试收款", "entry_token": "receipt-test-unique-token"}
    assert post(client, f"/documents/{identifier}/finance", receipt).status_code == 302
    assert post(client, f"/documents/{identifier}/finance", receipt).status_code == 302
    assert post(client, f"/documents/{identifier}/finance", {**receipt, "amount": "500"}).status_code == 409
    second = upload(client, text_pdf("TOKEN-CONFLICT"))
    assert save(client, second, amount="1000", currency="CNY").status_code == 302
    assert post(client, f"/documents/{second}/finance", receipt).status_code == 409
    assert post(client, f"/documents/{second}/archive").status_code == 302
    invoice = {"kind": "invoice", "amount": "600.00", "occurred_on": "2026-09-23", "reference": "TEST-INVOICE", "description": "验收测试开票", "entry_token": "invoice-test-unique-token"}
    assert post(client, f"/documents/{identifier}/finance", invoice).status_code == 302
    assert post(client, f"/documents/{identifier}/finance", {**invoice, "entry_token": "duplicate-invoice-another-token"}).status_code == 409
    content = client.get("/companies").get_data(as_text=True)
    assert "700.00" in content and "400.01" in content
    export = load_workbook(io.BytesIO(client.get("/companies/export.xlsx").data))
    assert export.sheetnames == ['公司账款汇总', '合同清单', '收款记录', '开票记录', '付款记录']
    assert export['公司账款汇总']['G2'].value == 700
    assert export['公司账款汇总']['I2'].value == 400.01
    with connect(app.config['DATA_DIR']) as conn:
        assert conn.execute('SELECT COUNT(*) FROM finance_records').fetchone()[0] == 2
        invoice_id = conn.execute("SELECT id FROM finance_records WHERE kind = 'invoice'").fetchone()[0]
    assert post(client, f"/finance/{invoice_id}/void", {'reason': '验收测试更正'}).status_code == 302
    assert post(client, f"/documents/{identifier}/finance", invoice).status_code == 409
    assert '1,000.01' in client.get('/companies').get_data(as_text=True)
    assert post(client, f"/documents/{identifier}/finance", {**receipt, 'amount': '-3', 'entry_token': 'negative-test-unique-token'}).status_code == 400


def test_missing_amount_currency_isolation_and_overpayment(app, client):
    first = upload(client, text_pdf('CNY'))
    second = upload(client, text_pdf('USD'))
    third = upload(client, text_pdf('UNKNOWN'))
    save(client, first, amount='100', currency='CNY')
    save(client, second, amount='200', currency='USD')
    save(client, third)
    post(client, f'/documents/{first}/finance', {'kind': 'receipt', 'amount': '120', 'occurred_on': '2026-09-23', 'entry_token': 'overpaid-test-unique-token'})
    assert post(client, f'/documents/{third}/finance', {'kind': 'receipt', 'amount': '1', 'occurred_on': '2026-09-23', 'entry_token': 'unknown-test-unique-token'}).status_code == 400
    report = load_workbook(io.BytesIO(client.get('/companies/export.xlsx').data))['公司账款汇总']
    assert report.max_row == 3
    values = {row[1]: row for row in report.iter_rows(min_row=2, values_only=True)}
    assert values['CNY'][3] == 1
    assert values['CNY'][6] == 0
    assert values['CNY'][9] == 20
    assert values['USD'][6] == 200
    assert 'TEST-ONLY' not in client.get('/?finance=unpaid').get_data(as_text=True)
    assert save(client, first, revision='2', amount='100', currency='USD').status_code == 400
