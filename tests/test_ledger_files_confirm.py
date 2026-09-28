import io
import zipfile

from openpyxl import load_workbook

from contractdb.db import connect
from test_workflows import app, client, post, save, text_pdf, upload  # noqa: F401


def finance(client, identifier, kind, amount, *, reference="", token, occurred_on="2026-09-24"):
    return post(client, f"/documents/{identifier}/finance", {
        "kind": kind, "amount": amount, "occurred_on": occurred_on,
        "reference": reference, "description": "台账测试", "entry_token": token,
    })


def test_payment_confirmation_snapshot_and_stale(app, client):
    identifier = upload(client)
    assert save(client, identifier, amount="1000", payable="600").status_code == 302
    assert finance(client, identifier, "receipt", "300", reference="R-001", token="confirm-receipt-0001").status_code == 302
    assert finance(client, identifier, "payment", "200", reference="P-001", token="confirm-payment-0001").status_code == 302
    assert post(client, f"/documents/{identifier}/confirm-payments", {"note": "已核对银行流水"}).status_code == 302
    with connect(app.config["DATA_DIR"]) as conn:
        row = conn.execute("SELECT * FROM payment_confirmations WHERE document_id = ?", (identifier,)).fetchone()
        assert row["received_minor"] == 30000
        assert row["unreceived_minor"] == 70000
        assert row["paid_minor"] == 20000
        assert row["unpaid_minor"] == 40000
        assert row["note"] == "已核对银行流水"
    detail = client.get(f"/documents/{identifier}").get_data(as_text=True)
    assert "款项确认" in detail and "已确认" in detail
    assert "对方公司已付" in detail
    assert finance(client, identifier, "receipt", "100", reference="R-002", token="confirm-receipt-0002").status_code == 302
    stale = client.get(f"/documents/{identifier}").get_data(as_text=True)
    assert "需重新确认" in stale
    blank = upload(client, text_pdf("BLANK"), name="未填金额.pdf")
    assert post(client, f"/documents/{blank}/confirm-payments", {}).status_code == 400
    assert post(client, f"/documents/{identifier}/confirm-payments", {"note": "x" * 501}).status_code == 400


def test_company_page_shows_confirmation_state(app, client):
    identifier = upload(client)
    assert save(client, identifier, amount="1000", payable="600").status_code == 302
    assert "未确认" in client.get("/companies").get_data(as_text=True)
    assert post(client, f"/documents/{identifier}/confirm-payments", {}).status_code == 302
    page = client.get("/companies").get_data(as_text=True)
    assert "已确认" in page and "已确认未收" in page


def test_ledger_page_filters_totals_and_void(app, client):
    identifier = upload(client)
    assert save(client, identifier, amount="1000", payable="600").status_code == 302
    assert finance(client, identifier, "receipt", "300", reference="R-001", token="ledger-receipt-0001").status_code == 302
    assert finance(client, identifier, "payment", "200", reference="P-001", token="ledger-payment-0001").status_code == 302
    assert finance(client, identifier, "invoice", "100", reference="INV-001", token="ledger-invoice-0001").status_code == 302
    page = client.get("/ledger").get_data(as_text=True)
    assert "款项台账" in page
    assert "我方收款" in page and "我方付款" in page and "我方开票" in page
    assert "300.00" in page and "200.00" in page and "100.00" in page
    payment_only = client.get("/ledger?kind=payment").get_data(as_text=True)
    assert "P-001" in payment_only and "R-001" not in payment_only
    assert client.get("/ledger?kind=bogus").status_code == 400
    assert client.get("/ledger?from=2026-10-01&to=2026-09-01").status_code == 400
    with connect(app.config["DATA_DIR"]) as conn:
        record = conn.execute("SELECT id FROM finance_records WHERE reference = 'R-001'").fetchone()[0]
    assert post(client, f"/finance/{record}/void", {"reason": "台账测试更正"}).status_code == 302
    voided = client.get("/ledger?status=void").get_data(as_text=True)
    assert "R-001" in voided and "已作废" in voided
    assert "R-001" not in client.get("/ledger?status=valid").get_data(as_text=True)


def test_ledger_export_workbook(app, client):
    identifier = upload(client)
    assert save(client, identifier, amount="1000", payable="600").status_code == 302
    assert finance(client, identifier, "receipt", "300", reference="R-001", token="export-receipt-0001").status_code == 302
    response = client.get("/ledger/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.sheetnames == ["收付款台账", "按币种汇总", "合同台账"]
    ledger_rows = list(workbook["收付款台账"].iter_rows(min_row=2, values_only=True))
    assert any(row[1] == "我方收款" and row[7] == 300 for row in ledger_rows)
    summary = {row[0]: row for row in workbook["按币种汇总"].iter_rows(min_row=2, values_only=True)}
    assert summary["CNY"][1] == 300 and summary["CNY"][4] == 1
    contracts = list(workbook["合同台账"].iter_rows(min_row=2, values_only=True))
    assert contracts[0][5] == 1000 and contracts[0][6] == 300 and contracts[0][7] == 700
    assert contracts[0][13] == "未确认"


def test_file_archive_page_and_bulk_actions(app, client):
    first = upload(client, text_pdf("FILE-A"), name="文件A.pdf")
    second = upload(client, text_pdf("FILE-B"), name="文件B.pdf")
    page = client.get("/files").get_data(as_text=True)
    assert "文件归档" in page and "文件A.pdf" in page and "文件B.pdf" in page
    assert client.get("/files?filetype=pdf").status_code == 200
    assert client.get("/files?status=archived").get_data(as_text=True).count("文件A.pdf") == 0
    assert post(client, "/files/bulk", {"action": "archive", "ids": [first, second]}).status_code == 302
    with connect(app.config["DATA_DIR"]) as conn:
        assert conn.execute("SELECT archived_at FROM documents WHERE id = ?", (first,)).fetchone()[0] is not None
        assert conn.execute("SELECT archived_at FROM documents WHERE id = ?", (second,)).fetchone()[0] is not None
    archived_page = client.get("/files?status=archived").get_data(as_text=True)
    assert "文件A.pdf" in archived_page and "文件B.pdf" in archived_page
    assert "文件A.pdf" not in client.get("/files?status=active").get_data(as_text=True)
    assert post(client, "/files/bulk", {"action": "archive", "ids": [first, second]}).status_code == 302
    assert post(client, "/files/bulk", {"action": "restore", "ids": [first]}).status_code == 302
    with connect(app.config["DATA_DIR"]) as conn:
        assert conn.execute("SELECT archived_at FROM documents WHERE id = ?", (first,)).fetchone()[0] is None
        assert conn.execute("SELECT archived_at FROM documents WHERE id = ?", (second,)).fetchone()[0] is not None


def test_file_archive_zip_contains_manifest_and_originals(app, client):
    upload(client, text_pdf("ZIP-A"), name="归档A.pdf")
    response = client.get("/files/export.zip")
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        names = archive.namelist()
        assert "文件归档清单.xlsx" in names
        assert "归档说明.txt" in names
        originals = [name for name in names if name.startswith("文件/")]
        assert any(name.endswith("归档A.pdf") for name in originals)
        manifest = load_workbook(io.BytesIO(archive.read("文件归档清单.xlsx")))
        values = list(manifest["文件归档清单"].iter_rows(min_row=2, values_only=True))
        assert values[0][2] == "归档A.pdf"
        assert len(values[0][12]) == 64


def test_file_and_ledger_permissions(app, client):
    anonymous = app.test_client()
    assert anonymous.get("/files").status_code == 302
    assert anonymous.get("/ledger").status_code == 302
    assert client.post("/files/bulk", data={"action": "archive", "ids": ["missing"]}).status_code == 400
    assert post(client, "/files/bulk", {"action": "archive"}).status_code == 400
    assert post(client, "/files/bulk", {"action": "bogus", "ids": ["missing"]}).status_code == 400
    assert client.get("/files?year=abcd").status_code == 400
    assert client.get("/files?status=bogus").status_code == 400
