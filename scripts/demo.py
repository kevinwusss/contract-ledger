"""Create fictional demo records through normal application routes."""
from __future__ import annotations

import getpass
import io
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from contractdb.web import create_app


def seed(root: Path, password: str):
    root = root.resolve()
    if len(password) < 10:
        raise ValueError("Choose a password of at least 10 characters.")
    if root.exists() and (not root.is_dir() or any(root.iterdir())):
        raise ValueError("Demo requires an empty data directory; existing data is never overwritten.")
    app = create_app({"DATA_DIR": root, "SYNC_EXTRACTION": True, "RESUME_EXTRACTION": False})
    client = app.test_client()
    client.get("/setup")

    def post(url, data):
        with client.session_transaction() as session:
            token = session["csrf"]
        response = client.post(url, data={"csrf_token": token, **data})
        if response.status_code != 302:
            raise RuntimeError(f"Demo step {url} failed with HTTP {response.status_code}")
        return response

    credentials = {"username": "demo-admin", "password": password}
    post("/setup", credentials)
    post("/login", credentials)
    identifiers = []
    for number, currency, amount, payable in [(1, "CNY", "10000", "2500"), (2, "USD", "2000", "0"), (3, "CNY", "", "")]:
        label = f"DEMO-{number:03d}"
        document = Document()
        document.add_heading(f"Synthetic demonstration: {label}", 0)
        document.add_paragraph("Fictional software fixture. Not a legal contract or real transaction.")
        document.add_paragraph(f"Reference: {label}; project: Example Service {number}.")
        if amount:
            document.add_paragraph(f"Fictional receivable: {currency} {amount}; payable: {currency} {payable}.")
        else:
            document.add_paragraph("Amounts and counterparty intentionally unspecified for review.")
        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)
        identifier = post("/upload", {"files": (buffer, f"{label}.docx")}).location.rsplit("/", 1)[-1]
        identifiers.append(identifier)
        if number == 3:
            continue
        post(f"/documents/{identifier}/save", {
            "revision": "1", "title": f"Synthetic Service Contract {number}",
            "company": f"Example Client {'A' if number == 1 else 'B'} (Fictional)",
            "project": f"Example Service {number}", "contract_number": label,
            "signed_date": "2026-09-01", "verified": "1", "type_name": "Demo Service",
            "subtype_name": "Synthetic", "amount": amount, "payable": payable,
            "currency": currency, "notes": "Synthetic demo only; never operational evidence.",
        })
        entries = [("receipt", "4000"), ("payment", "1000"), ("invoice", "3000")] if number == 1 else [("receipt", "500")]
        for kind, value in entries:
            post(f"/documents/{identifier}/finance", {
                "kind": kind, "amount": value, "occurred_on": "2026-09-02",
                "reference": f"{label}-{kind}", "description": "Fictional demo transaction",
                "entry_token": f"synthetic-{label}-{kind}-entry",
            })
        if number == 1:
            post(f"/documents/{identifier}/confirm-payments", {"note": "Synthetic initial confirmation"})
    return app, identifiers


def main():
    if os.environ.get("CONTRACTDB_DATA_DIR"):
        raise SystemExit("Unset CONTRACTDB_DATA_DIR first. This demo uses only this copy's data directory.")
    target = ROOT / "data"
    if target.exists() and (not target.is_dir() or any(target.iterdir())):
        raise SystemExit("Data already exists. Use a fresh repository copy for the demo.")
    password = getpass.getpass("New demo-admin password (at least 10 characters): ")
    if password != getpass.getpass("Repeat password: "):
        raise SystemExit("Passwords do not match.")
    seed(target, password)
    print("Created 3 fictional contracts. Log in as demo-admin with your chosen password.")
    print("Start: python -m contractdb serve --host 127.0.0.1 --port 8000")


if __name__ == "__main__":
    main()
