from __future__ import annotations

import hmac
import io
import ipaddress
import json
import math
import re
import secrets
import socket
import sqlite3
import tempfile
import subprocess
import threading
import time
import zipfile
from datetime import date, datetime, timedelta
from functools import wraps
from pathlib import Path
from urllib.parse import urlencode

from argon2 import PasswordHasher
from argon2.exceptions import InvalidHashError, VerificationError
from flask import Flask, abort, flash, g, jsonify, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import db
from .backup import create_backup
from .config import MAX_UPLOAD_BYTES, PROJECT_ROOT, data_dir, ensure_data_dirs
from .service import DuplicateError, ExtractionQueue, UploadError, ingest
from .finance import CURRENCIES, CONFIRMATION_LABELS, amount_summary, balance_snapshot, company_summary, confirmation_status, parse_money


password_hasher = PasswordHasher()
EVENT_NAMES = {"import": "导入原件", "extract": "完成信息提取", "extract_error": "提取失败", "edit": "保存信息", "archive": "归档", "restore": "恢复", "export_excel": "导出 Excel", "export_zip": "导出原件 ZIP", "user_create": "创建账号", "user_update": "修改账号", "taxonomy": "维护分类", "backup": "创建备份", "setup": "创建管理员"}
FIELD_NAMES = {"company": "合同相对方", "project": "定点项目", "contract_number": "合同编号", "signed_date": "签订日期", "subtype": "子类别", "amount_hint": "原件金额（收付方向待核对）"}
EVENT_NAMES.update(finance_add="登记收款 / 开票", finance_void="作废收款 / 开票记录")
EVENT_NAMES.update(payment_confirm="确认款项", ledger_export="导出收付款台账", file_archive_export="导出文件归档包",
                   file_archive_batch="批量归档文件", file_restore_batch="批量恢复文件")


def create_app(config: dict | None = None) -> Flask:
    app = Flask(__name__)
    app.config.update(DATA_DIR=data_dir(), MAX_CONTENT_LENGTH=MAX_UPLOAD_BYTES + 1024 * 1024,
                      SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
                      PERMANENT_SESSION_LIFETIME=timedelta(hours=12), SYNC_EXTRACTION=False,
                      RESUME_EXTRACTION=True)
    if config:
        app.config.update(config)
    root = Path(app.config["DATA_DIR"]).resolve()
    app.config["DATA_DIR"] = root
    ensure_data_dirs(root)
    db.init_db(root)
    secret_path = root / "session.key"
    try:
        with secret_path.open("x", encoding="ascii") as output:
            output.write(secrets.token_hex(32))
    except FileExistsError:
        pass
    app.secret_key = secret_path.read_text(encoding="ascii").strip()
    queue = ExtractionQueue(root, synchronous=app.config["SYNC_EXTRACTION"])
    app.extensions["extraction_queue"] = queue
    if app.config["RESUME_EXTRACTION"]:
        queue.resume()
    attempts: dict[str, list[float]] = {}
    attempts_lock = threading.Lock()

    def connection():
        if "db" not in g:
            g.db = db.connect(root)
        return g.db

    @app.teardown_appcontext
    def close_connection(_error):
        conn = g.pop("db", None)
        if conn:
            conn.close()

    @app.before_request
    def load_user_and_csrf():
        g.user = None
        if session.get("user_id"):
            user = connection().execute("SELECT * FROM users WHERE id = ? AND active = 1", (session["user_id"],)).fetchone()
            if user and user["session_version"] == session.get("session_version"):
                g.user = user
            else:
                session.clear()
        session.setdefault("csrf", secrets.token_urlsafe(32))
        if request.method == "POST":
            submitted = request.form.get("csrf_token", "")
            if not hmac.compare_digest(submitted, session["csrf"]):
                abort(400, "页面已过期，请刷新后重新提交。")

    @app.after_request
    def headers(response):
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Referrer-Policy"] = "same-origin"
        response.headers["X-Frame-Options"] = "SAMEORIGIN"
        response.headers["Content-Security-Policy"] = "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; frame-src 'self'; object-src 'none'; base-uri 'self'; form-action 'self'; frame-ancestors 'self'"
        if request.endpoint != "static":
            response.headers["Cache-Control"] = "no-store"
        return response

    @app.context_processor
    def common_context():
        lan_url = None
        try:
            addresses = {item[4][0] for item in socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET)}
            address = next((value for value in sorted(addresses)
                            if ipaddress.ip_address(value).is_private
                            and not ipaddress.ip_address(value).is_loopback
                            and not ipaddress.ip_address(value).is_link_local
                            and not ipaddress.ip_address(value).is_unspecified), None)
            if address:
                lan_url = f"http://{address}:{app.config.get('SERVICE_PORT', 8000)}"
        except OSError:
            pass
        return {"csrf_token": session.get("csrf", ""), "event_names": EVENT_NAMES,
                "field_names": FIELD_NAMES, "current_user": g.user, "currencies": CURRENCIES,
                "lan_url": lan_url}

    @app.template_filter("money")
    def money(value):
        if value is None:
            return "待录入"
        from decimal import Decimal
        return f"{Decimal(value) / 100:,.2f}"

    @app.template_filter("money_input")
    def money_input(value):
        if value is None:
            return ""
        from decimal import Decimal
        return f"{Decimal(value) / 100:.2f}"

    @app.template_filter("localtime")
    def localtime(value):
        if not value:
            return "—"
        from datetime import datetime
        return datetime.fromisoformat(value).astimezone().strftime("%Y-%m-%d %H:%M")

    @app.template_filter("filesize")
    def filesize(value):
        return f"{value / 1024 / 1024:.1f} MB" if value >= 1024 * 1024 else f"{value / 1024:.0f} KB"

    def login_required(function):
        @wraps(function)
        def wrapped(*args, **kwargs):
            if not g.user:
                return redirect(url_for("login"))
            return function(*args, **kwargs)
        return wrapped

    def admin_required(function):
        @wraps(function)
        @login_required
        def wrapped(*args, **kwargs):
            if g.user["role"] != "admin":
                abort(403, "此操作需要管理员权限。")
            return function(*args, **kwargs)
        return wrapped

    def document(identifier):
        row = connection().execute(db.document_detail_query(), (identifier,)).fetchone()
        if not row:
            abort(404, "未找到这份合同。")
        return row

    def filters_from_request():
        filters = {key: request.args.get(key, "").strip()[:200] for key in ("q", "company", "type", "subtype", "project", "status", "from", "to", "finance", "currency", "sort")}
        for key in ("company", "type", "subtype", "project"):
            if filters[key] and (not filters[key].isdigit() or len(filters[key]) > 18):
                abort(400, "筛选条件不正确，请清空后重新选择。")
        for key in ("from", "to"):
            if filters[key]:
                try:
                    date.fromisoformat(filters[key])
                except ValueError:
                    abort(400, "日期格式不正确。")
        if filters['currency'] and filters['currency'] not in CURRENCIES:
            abort(400, '请选择支持的币种。')
        if filters['from'] and filters['to'] and filters['from'] > filters['to']:
            abort(400, '开始日期不能晚于结束日期。')
        return filters

    def filtered_rows():
        filters = filters_from_request()
        sql, parameters = db.document_query(filters, archived=request.args.get("archived") == "1")
        return connection().execute(sql, parameters).fetchall()

    @app.route("/setup", methods=["GET", "POST"])
    def setup():
        if connection().execute("SELECT 1 FROM users LIMIT 1").fetchone():
            return redirect(url_for("login"))
        if request.remote_addr not in {"127.0.0.1", "::1"}:
            abort(403, "请在服务主机上访问 127.0.0.1 创建首个管理员。")
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            if not re.fullmatch(r"[\w\-.]{2,40}", username) or len(password) < 10:
                flash("账号需为 2–40 位文字、数字或下划线；密码至少 10 位。", "error")
            else:
                with db.transaction(root) as conn:
                    if conn.execute("SELECT 1 FROM users LIMIT 1").fetchone():
                        abort(409, "管理员已创建，请登录。")
                    cursor = conn.execute("INSERT INTO users(username, password_hash, role, created_at) VALUES (?, ?, 'admin', ?)", (username, password_hasher.hash(password), db.now()))
                    db.audit(conn, "setup", user_id=cursor.lastrowid)
                flash("管理员已创建，请登录。", "success")
                return redirect(url_for("login"))
        return render_template("login.html", setup=True)

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if not connection().execute("SELECT 1 FROM users LIMIT 1").fetchone():
            return redirect(url_for("setup"))
        if g.user:
            return redirect(url_for("index"))
        if request.method == "POST":
            key = request.remote_addr or "unknown"
            with attempts_lock:
                recent = [item for item in attempts.get(key, []) if item > time.monotonic() - 300]
                attempts[key] = recent
                if len(recent) >= 10:
                    abort(429, "登录尝试过于频繁，请 5 分钟后重试。")
            user = connection().execute("SELECT * FROM users WHERE username = ? AND active = 1", (request.form.get("username", "").strip(),)).fetchone()
            valid = False
            if user:
                try:
                    valid = password_hasher.verify(user["password_hash"], request.form.get("password", ""))
                except (VerificationError, InvalidHashError):
                    pass
            if valid:
                session.clear()
                session.update(user_id=user["id"], session_version=user["session_version"], csrf=secrets.token_urlsafe(32))
                session.permanent = True
                with attempts_lock:
                    attempts.pop(key, None)
                return redirect(url_for("index"))
            with attempts_lock:
                attempts[key].append(time.monotonic())
            flash("账号或密码不正确，或账号已停用。", "error")
        return render_template("login.html", setup=False)

    @app.post("/logout")
    def logout():
        session.clear()
        return redirect(url_for("login"))

    @app.get("/")
    @login_required
    def index():
        filters = filters_from_request()
        archived = request.args.get("archived") == "1"
        sql, parameters = db.document_query(filters, archived=archived)
        total = connection().execute(f"SELECT COUNT(*) FROM ({sql})", parameters).fetchone()[0]
        page = max(1, request.args.get("page", 1, type=int) or 1)
        pages = max(1, math.ceil(total / 25))
        page = min(page, pages)
        rows = connection().execute(sql + " LIMIT 25 OFFSET ?", parameters + [(page - 1) * 25]).fetchall()
        summary = amount_summary(connection().execute(sql, parameters))
        pending_count = connection().execute("SELECT COUNT(*) FROM documents WHERE review_status = 'pending' AND archived_at IS NULL").fetchone()[0]
        query = {key: value for key, value in filters.items() if value}
        if archived:
            query["archived"] = "1"
        finance_links = {key: url_for('index') + '?' + urlencode({**query, 'finance': key}) for key in ('unpaid', 'payable', 'uninvoiced', 'amount_missing', 'payable_missing', 'exception')}
        return render_template("index.html", rows=rows, filters=filters, archived=archived,
                               total=total, page=page, pages=pages, pending_count=pending_count, amount_groups=summary, finance_links=finance_links,
                               query=urlencode(query), previous=urlencode({**query, "page": page - 1}),
                               following=urlencode({**query, "page": page + 1}), **db.choices(connection()))

    @app.route("/upload", methods=["GET", "POST"])
    @login_required
    def upload():
        if request.method == "POST":
            files = [item for item in request.files.getlist("files") if item.filename]
            if not files:
                flash("请先选择要导入的合同文件。", "error")
                return redirect(url_for("upload"))
            if len(files) > 20:
                abort(400, "一次最多导入 20 个文件。")
            identifiers = []
            for file in files:
                try:
                    identifier = ingest(root, file.stream, file.filename, user_id=g.user["id"])
                    queue.submit(identifier)
                    identifiers.append(identifier)
                except DuplicateError as exc:
                    flash(f"{file.filename} 已存在，未重复导入。", "info")
                    if len(files) == 1:
                        return redirect(url_for("detail", identifier=exc.document_id))
                except UploadError as exc:
                    flash(f"{file.filename}：{exc}", "error")
            if identifiers:
                flash(f"已导入 {len(identifiers)} 份文件。核对信息后可标记为已核对。", "success")
                return redirect(url_for("detail", identifier=identifiers[0]) if len(identifiers) == 1 else url_for("index", status="pending"))
        return render_template("upload.html")

    @app.get("/documents/<identifier>")
    @login_required
    def detail(identifier):
        row = document(identifier)
        events = connection().execute("SELECT a.*, u.username FROM audit_events a LEFT JOIN users u ON u.id = a.user_id WHERE document_id = ? ORDER BY a.id DESC LIMIT 30", (identifier,)).fetchall()
        records = connection().execute("SELECT f.*, u.username FROM finance_records f LEFT JOIN users u ON u.id = f.created_by WHERE document_id = ? ORDER BY occurred_on DESC, f.id DESC", (identifier,)).fetchall()
        received = sum(r["amount_minor"] for r in records if r["kind"] == "receipt" and not r["voided_at"])
        invoiced = sum(r["amount_minor"] for r in records if r["kind"] == "invoice" and not r["voided_at"])
        paid = sum(r["amount_minor"] for r in records if r["kind"] == "payment" and not r["voided_at"])
        confirmations = connection().execute(
            "SELECT p.*, u.username FROM payment_confirmations p LEFT JOIN users u ON u.id = p.created_by "
            "WHERE p.document_id = ? ORDER BY p.id DESC LIMIT 20", (identifier,)).fetchall()
        return render_template("detail.html", doc=row, candidates=json.loads(row["candidates_json"]),
                               events=events, records=records, received=received, invoiced=invoiced, paid=paid,
                               confirmations=confirmations, confirmation=confirmations[0] if confirmations else None,
                               confirmation_state=confirmation_status(row), confirmation_labels=CONFIRMATION_LABELS,
                               entry_token=secrets.token_urlsafe(24), today=date.today().isoformat(), **db.choices(connection()))

    @app.post("/documents/<identifier>/save")
    @login_required
    def save_document(identifier):
        current = document(identifier)
        if current["archived_at"]:
            abort(400, "请先恢复归档合同，再修改信息。")
        fields = {key: request.form.get(key, "").strip() for key in ("title", "company", "project", "contract_number", "signed_date", "notes")}
        for key, limit in (("title", 250), ("company", 200), ("project", 200), ("contract_number", 120), ("notes", 5000)):
            if len(fields[key]) > limit:
                abort(400, f"字段过长：{FIELD_NAMES.get(key, key)}")
        if fields["signed_date"]:
            try:
                date.fromisoformat(fields["signed_date"])
            except ValueError:
                abort(400, "签订日期不正确。")
        type_id = request.form.get("type_id", "") or None
        subtype_id = request.form.get("subtype_id", "") or None
        type_name = request.form.get("type_name", "").strip()
        subtype_name = request.form.get("subtype_name", "").strip()
        if len(type_name) > 200 or len(subtype_name) > 200:
            abort(400, "类型和子类别最多填写 200 字。")
        if "type_name" in request.form:
            type_id = subtype_id = None
            if subtype_name and not type_name:
                abort(400, "填写子类别时，请同时填写合同类型。")
        try:
            amount_minor = parse_money(request.form.get("amount", ""), optional=True)
            payable_minor = parse_money(request.form.get("payable", ""), optional=True) if 'payable' in request.form else current['payable_minor']
        except ValueError as exc:
            abort(400, str(exc))
        currency = request.form.get("currency", "CNY")
        if currency not in CURRENCIES:
            abort(400, "请选择支持的币种。")
        with db.transaction(root) as conn:
            latest = conn.execute("SELECT revision, archived_at FROM documents WHERE id = ?", (identifier,)).fetchone()
            if latest["revision"] != request.form.get("revision", type=int) or latest["archived_at"]:
                abort(409, "这份合同已被其他操作更新。请返回详情页刷新，核对最新内容后重新保存。")
            has_finance = conn.execute("SELECT 1 FROM finance_records WHERE document_id = ? AND voided_at IS NULL LIMIT 1", (identifier,)).fetchone()
            has_receivable = conn.execute("SELECT 1 FROM finance_records WHERE document_id = ? AND kind IN ('receipt', 'invoice') AND voided_at IS NULL", (identifier,)).fetchone()
            has_payment = conn.execute("SELECT 1 FROM finance_records WHERE document_id = ? AND kind = 'payment' AND voided_at IS NULL", (identifier,)).fetchone()
            if (has_receivable and amount_minor is None) or (has_payment and payable_minor is None) or (has_finance and (not fields['company'] or currency != current['currency'])):
                abort(400, "已有有效收款或开票记录时，不能清空合同金额、公司或更改币种。请先核对相关记录。")
            if type_name:
                conn.execute("INSERT OR IGNORE INTO import_types(name) VALUES (?)", (type_name,))
                type_id = conn.execute("SELECT id FROM import_types WHERE name = ?", (type_name,)).fetchone()[0]
                if subtype_name:
                    conn.execute("INSERT OR IGNORE INTO subtypes(type_id, name) VALUES (?, ?)", (type_id, subtype_name))
                    subtype_id = conn.execute("SELECT id FROM subtypes WHERE type_id = ? AND name = ?", (type_id, subtype_name)).fetchone()[0]
            if type_id and not conn.execute("SELECT 1 FROM import_types WHERE id = ?", (type_id,)).fetchone():
                abort(400, "导入类型不存在。")
            if subtype_id and not conn.execute("SELECT 1 FROM subtypes WHERE id = ? AND type_id = ?", (subtype_id, type_id)).fetchone():
                abort(400, "子类别与导入类型不匹配。")
            company_id = None
            project_id = None
            if fields["company"]:
                conn.execute("INSERT OR IGNORE INTO companies(name) VALUES (?)", (fields["company"],))
                company_id = conn.execute("SELECT id FROM companies WHERE name = ?", (fields["company"],)).fetchone()[0]
            if fields["project"]:
                conn.execute("INSERT OR IGNORE INTO projects(company_id, name) VALUES (?, ?)", (company_id, fields["project"]))
                project_id = conn.execute("SELECT id FROM projects WHERE company_id IS ? AND name = ?", (company_id, fields["project"])).fetchone()[0]
            status = "verified" if request.form.get("verified") == "1" else "pending"
            updates = {"title": fields["title"] or None, "contract_number": fields["contract_number"] or None,
                       "signed_date": fields["signed_date"] or None, "company_id": company_id, "type_id": type_id,
                       "subtype_id": subtype_id, "project_id": project_id, "notes": fields["notes"] or None,
                       "review_status": status, "amount_minor": amount_minor, "payable_minor": payable_minor, "currency": currency}
            conn.execute("UPDATE documents SET " + ", ".join(f"{key} = ?" for key in updates) + ", updated_at = ?, updated_by = ?, revision = revision + 1 WHERE id = ?",
                         list(updates.values()) + [db.now(), g.user["id"], identifier])
            changes = {key: {"before": current[key], "after": value} for key, value in updates.items() if str(current[key] or "") != str(value or "")}
            db.audit(conn, "edit", document_id=identifier, user_id=g.user["id"], details={"changes": changes})
        flash("合同信息已保存。" + ("已标记为已核对。" if status == "verified" else "仍为待核对状态。"), "success")
        return redirect(url_for("detail", identifier=identifier))

    @app.post("/documents/<identifier>/finance")
    @login_required
    def add_finance(identifier):
        kind = request.form.get("kind")
        if kind not in {"receipt", "invoice", "payment"}:
            abort(400, "记录类型不正确。")
        try:
            amount = parse_money(request.form.get("amount", ""), positive=True)
            occurred_on = date.fromisoformat(request.form.get("occurred_on", "")).isoformat()
        except ValueError as exc:
            abort(400, f"金额或日期不正确：{exc}")
        reference = request.form.get("reference", "").strip()
        description = request.form.get("description", "").strip()
        token = request.form.get("entry_token", "")
        if len(reference) > 120 or len(description) > 2000 or not 16 <= len(token) <= 80:
            abort(400, "编号或说明过长，或提交标识无效。请刷新后重试。")
        if kind == "invoice" and not reference:
            abort(400, "请填写发票号码，便于核对和防止重复登记。")
        with db.transaction(root) as conn:
            doc = conn.execute("SELECT * FROM documents WHERE id = ?", (identifier,)).fetchone()
            if not doc:
                abort(404)
            total_key = 'payable_minor' if kind == 'payment' else 'amount_minor'
            if doc["archived_at"] or doc[total_key] is None or not doc["company_id"]:
                abort(400, "请先填写公司和对应方向的应收或应付总额并保存；归档合同需先恢复。")
            existing = conn.execute("SELECT * FROM finance_records WHERE entry_token = ?", (token,)).fetchone()
            if existing:
                submitted = {"document_id": identifier, "kind": kind, "amount_minor": amount,
                             "occurred_on": occurred_on, "reference": reference, "description": description}
                if existing["voided_at"] or any(existing[key] != value for key, value in submitted.items()):
                    abort(409, "提交标识已用于其他内容或已作废记录。请刷新页面，核对明细后重新登记。")
                flash("这笔记录已保存，未重复登记。", "info")
                return redirect(url_for("detail", identifier=identifier) + "#finance")
            if kind == "invoice" and conn.execute("SELECT 1 FROM finance_records WHERE document_id = ? AND kind = 'invoice' AND reference = ? AND voided_at IS NULL", (identifier, reference)).fetchone():
                abort(409, "这份合同已登记相同发票号码，请核对后再录入。")
            conn.execute("INSERT INTO finance_records(document_id, kind, amount_minor, occurred_on, reference, description, entry_token, created_by, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                         (identifier, kind, amount, occurred_on, reference, description, token, g.user["id"], db.now()))
            db.audit(conn, "finance_add", document_id=identifier, user_id=g.user["id"], details={"kind": kind, "amount_minor": amount, "currency": doc["currency"], "reference": reference, "occurred_on": occurred_on})
        flash({'receipt': '收款记录已保存。', 'payment': '付款记录已保存。', 'invoice': '开票记录已保存。'}[kind], "success")
        return redirect(url_for("detail", identifier=identifier) + "#finance")

    @app.post("/finance/<int:record_id>/void")
    @login_required
    def void_finance(record_id):
        reason = request.form.get("reason", "").strip()
        if not reason or len(reason) > 500:
            abort(400, "请填写作废原因（最多 500 字）。")
        with db.transaction(root) as conn:
            record = conn.execute("SELECT f.*, d.archived_at FROM finance_records f JOIN documents d ON d.id = f.document_id WHERE f.id = ?", (record_id,)).fetchone()
            if not record:
                abort(404)
            if record["archived_at"]:
                abort(400, "请先恢复归档合同再修改财务记录。")
            conn.execute("UPDATE finance_records SET voided_at = ?, void_reason = ? WHERE id = ? AND voided_at IS NULL", (db.now(), reason, record_id))
            db.audit(conn, "finance_void", document_id=record["document_id"], user_id=g.user["id"], details={"record_id": record_id, "reason": reason})
        flash("记录已作废，金额已重新计算。原记录仍保留供核对。", "success")
        return redirect(url_for("detail", identifier=record["document_id"]) + "#finance")

    @app.post("/documents/<identifier>/confirm-payments")
    @login_required
    def confirm_payments(identifier):
        note = request.form.get("note", "").strip()
        if len(note) > 500:
            abort(400, "确认备注最多 500 字。")
        with db.transaction(root) as conn:
            row = conn.execute(db.document_detail_query(), (identifier,)).fetchone()
            if not row:
                abort(404)
            if row["archived_at"]:
                abort(400, "归档合同需先恢复，再确认款项。")
            if row["amount_minor"] is None and row["payable_minor"] is None:
                abort(400, "请先填写应收或应付总额，再确认款项。")
            snapshot = balance_snapshot(row)
            conn.execute(
                """INSERT INTO payment_confirmations(document_id, currency, amount_minor, received_minor,
                   unreceived_minor, payable_minor, paid_minor, unpaid_minor, invoiced_minor, uninvoiced_minor,
                   note, created_by, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (identifier, snapshot["currency"], snapshot["amount_minor"], snapshot["received_minor"],
                 snapshot["unreceived_minor"], snapshot["payable_minor"], snapshot["paid_minor"],
                 snapshot["unpaid_minor"], snapshot["invoiced_minor"], snapshot["uninvoiced_minor"],
                 note, g.user["id"], db.now()),
            )
            db.audit(conn, "payment_confirm", document_id=identifier, user_id=g.user["id"],
                     details={**snapshot, "note": note})
        flash("款项已确认。之后的收款、付款或开票变动会自动提示重新确认。", "success")
        return redirect(url_for("detail", identifier=identifier) + "#confirmation")

    @app.get("/companies")
    @login_required
    def company_accounts():
        filters = filters_from_request()
        sql, params = db.document_query(filters)
        rows = connection().execute(sql, params).fetchall()
        open_items = [row for row in rows if row["company_id"] and (
            (row["amount_minor"] is not None and
             (row["amount_minor"] > row["received_minor"] or row["amount_minor"] > row["invoiced_minor"]))
            or (row["payable_minor"] is not None and row["payable_minor"] > row["paid_minor"]))]
        open_items.sort(key=lambda row: (row["company_name"], row["currency"], row["title"] or row["original_filename"]))
        return render_template("companies.html", groups=company_summary(rows), filters=filters,
                               open_items=open_items,
                               missing_company=sum(1 for row in rows if not row["company_id"]),
                               query=urlencode({key: value for key, value in filters.items() if value}), **db.choices(connection()))

    @app.post("/documents/<identifier>/archive")
    @login_required
    def archive_document(identifier):
        document(identifier)
        with db.transaction(root) as conn:
            conn.execute("UPDATE documents SET archived_at = ?, updated_at = ?, revision = revision + 1 WHERE id = ? AND archived_at IS NULL", (db.now(), db.now(), identifier))
            db.audit(conn, "archive", document_id=identifier, user_id=g.user["id"])
        flash("合同已归档，原件保留，可从归档库恢复。", "success")
        return redirect(url_for("index"))

    @app.post("/documents/<identifier>/restore")
    @login_required
    def restore_document(identifier):
        row = document(identifier)
        try:
            with db.transaction(root) as conn:
                conn.execute("UPDATE documents SET archived_at = NULL, updated_at = ?, revision = revision + 1 WHERE id = ?", (db.now(), identifier))
                db.audit(conn, "restore", document_id=identifier, user_id=g.user["id"])
        except sqlite3.IntegrityError:
            flash("合同库中已有相同原件，无法重复恢复。请先检查该文件。", "error")
            return redirect(url_for("detail", identifier=row["id"]))
        flash("合同已恢复。", "success")
        return redirect(url_for("detail", identifier=identifier))

    @app.post("/documents/<identifier>/extract")
    @login_required
    def retry_extraction(identifier):
        row = document(identifier)
        if row["extraction_status"] in {"queued", "processing"}:
            flash("提取正在进行，请稍候。", "info")
        else:
            with db.transaction(root) as conn:
                conn.execute("UPDATE documents SET extraction_status = 'queued' WHERE id = ?", (identifier,))
            queue.submit(identifier)
            flash("已重新提交本地提取。", "info")
        return redirect(url_for("detail", identifier=identifier))

    @app.get("/api/documents/<identifier>/status")
    @login_required
    def extraction_status(identifier):
        row = document(identifier)
        return jsonify(status=row["extraction_status"])

    @app.get("/documents/<identifier>/file")
    @login_required
    def original_file(identifier):
        row = document(identifier)
        inline = request.args.get("inline") == "1" and row["mime_type"] in {"application/pdf", "image/png", "image/jpeg"}
        path = root / "originals" / row["stored_filename"]
        if not path.is_file():
            abort(404, "原件文件缺失，请联系管理员从备份恢复。")
        return send_file(path, mimetype=row["mime_type"], as_attachment=not inline,
                         download_name=row["original_filename"], conditional=True)

    def excel_content(rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "合同清单"
        columns = [("公司名称", "company_name"), ("导入类型", "type_name"), ("子类别", "subtype_name"),
                   ("定点项目", "project_name"), ("合同名称", "title"), ("合同编号", "contract_number"),
                   ("签订日期", "signed_date"), ("核对状态", "review_status"), ("备注", "notes"),
                   ("原始文件名", "original_filename"), ("上传时间", "created_at"), ("更新时间", "updated_at"),
                   ("归档时间", "archived_at"), ("记录编号", "id"), ("原件 SHA256", "sha256")]
        columns += [("币种", "currency"), ("对方应付我方总额", "amount_minor"), ("已收款", "received_minor"),
                    ("未收款", "unpaid_minor"), ("已开票", "invoiced_minor"), ("未开票", "uninvoiced_minor"),
                    ("多收款", "overpaid_minor"), ("超开票", "overinvoiced_minor")]
        columns += [("我方应付总额", "payable_minor"), ("我方已付款", "paid_minor"), ("我方未付款", "outstanding_payable_minor")]
        sheet.append([label for label, _ in columns])
        for row in rows:
            values = []
            for _, key in columns:
                if key == "outstanding_payable_minor":
                    value = max(row["payable_minor"] - row["paid_minor"], 0) if row["payable_minor"] is not None else None
                elif key in {"unpaid_minor", "uninvoiced_minor", "overpaid_minor", "overinvoiced_minor"}:
                    if row["amount_minor"] is None:
                        value = None
                    else:
                        paid = row["received_minor"] if key in {"unpaid_minor", "overpaid_minor"} else row["invoiced_minor"]
                        value = max(0, row["amount_minor"] - paid) if key.startswith("un") else max(0, paid - row["amount_minor"])
                else:
                    value = row[key]
                if key == "review_status":
                    value = "已核对" if value == "verified" else "待核对"
                if key.endswith("_at") and value:
                    value = localtime(value)
                if key.endswith("_minor"):
                    values.append(value / 100 if value is not None else None)
                else:
                    values.append(value or "")
            sheet.append(values)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="183B39")
            cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00;[Red](#,##0.00);0.00'
                else:
                    cell.data_type = "s"
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        widths = [32, 15, 18, 30, 38, 25, 15, 13, 38, 55, 22, 22, 22, 36, 68]
        from openpyxl.utils import get_column_letter
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index in range(16, len(columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 20
        ids = {row["id"] for row in rows}
        row_map = {row["id"]: row for row in rows}
        for kind, label in (("receipt", "收款记录"), ("invoice", "开票记录"), ("payment", "付款记录")):
            detail_sheet = workbook.create_sheet(label)
            detail_sheet.append(["公司名称", "合同名称 / 原文件名", "定点项目", "币种", "金额", "日期", "凭证 / 发票号", "具体说明", "记录状态", "作废原因"])
            records = connection().execute("SELECT * FROM finance_records WHERE kind = ? ORDER BY occurred_on, id", (kind,)).fetchall()
            for record in records:
                if record["document_id"] not in ids:
                    continue
                doc = row_map[record["document_id"]]
                detail_sheet.append([doc["company_name"], doc["title"] or doc["original_filename"], doc["project_name"], doc["currency"], record["amount_minor"] / 100, record["occurred_on"], record["reference"], record["description"], "已作废" if record["voided_at"] else "有效", record["void_reason"]])
            for cells in detail_sheet:
                for cell in cells:
                    if not isinstance(cell.value, (int, float)):
                        cell.data_type = 's'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    cell.font = Font(name='Microsoft YaHei', size=10)
            for cell in detail_sheet[1]:
                cell.fill = PatternFill('solid', fgColor='183B39')
                cell.font = Font(name='Microsoft YaHei', color='FFFFFF', bold=True)
            for column in range(1, 11):
                detail_sheet.column_dimensions[get_column_letter(column)].width = 28 if column != 2 else 48
            detail_sheet.freeze_panes = 'A2'
            detail_sheet.auto_filter.ref = detail_sheet.dimensions
        if request.endpoint == 'export_company_accounts':
            summary = workbook.create_sheet('公司账款汇总', 0)
            summary.append(['公司', '币种', '合同份数', '金额未录入份数', '合同含税总额', '已收款', '未收款', '已开票', '未开票', '多收款', '超开票', '我方应付总额', '我方已付款', '我方未付款', '我方超额付款', '应付金额待录入份数'])
            for group in company_summary(rows):
                summary.append([group['company_name'], group['currency'], group['contract_count'], group['unknown_count']] + [group[key] / 100 for key in ('amount_minor', 'received_minor', 'unpaid_minor', 'invoiced_minor', 'uninvoiced_minor', 'overpaid_minor', 'overinvoiced_minor', 'payable_minor', 'paid_minor', 'outstanding_payable_minor', 'excess_payment_minor')] + [group['payable_unknown_count']])
            for cells in summary:
                for cell in cells:
                    if not isinstance(cell.value, (int, float)):
                        cell.data_type = 's'
                    cell.font = Font(name='Microsoft YaHei', size=10)
            for cell in summary[1]:
                cell.fill = PatternFill('solid', fgColor='183B39')
                cell.font = Font(name='Microsoft YaHei', color='FFFFFF', bold=True)
            for column in range(1, 17):
                summary.column_dimensions[get_column_letter(column)].width = 23 if column != 1 else 38
            summary.freeze_panes = 'A2'
            summary.auto_filter.ref = summary.dimensions
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    @app.get("/export.xlsx")
    @login_required
    def export_excel():
        rows = filtered_rows()
        with db.transaction(root) as conn:
            db.audit(conn, "export_excel", user_id=g.user["id"], details={"count": len(rows), "filters": dict(request.args)})
        return send_file(excel_content(rows), as_attachment=True, download_name="合同清单.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.get('/companies/export.xlsx')
    @login_required
    def export_company_accounts():
        rows = filtered_rows()
        with db.transaction(root) as conn:
            db.audit(conn, 'export_excel', user_id=g.user['id'], details={'report': 'company_accounts', 'count': len(rows), 'filters': dict(request.args)})
        return send_file(excel_content(rows), as_attachment=True, download_name='公司账款与开票明细.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def ledger_filters_from_request():
        filters = {key: request.args.get(key, "").strip()[:200] for key in ("q", "company", "kind", "currency", "from", "to", "status")}
        if filters["company"] and (not filters["company"].isdigit() or len(filters["company"]) > 18):
            abort(400, "公司筛选不正确。")
        if filters["kind"] and filters["kind"] not in {"receipt", "payment", "invoice"}:
            abort(400, "记录类型不正确。")
        if filters["currency"] and filters["currency"] not in CURRENCIES:
            abort(400, "请选择支持的币种。")
        if filters["status"] and filters["status"] not in {"valid", "void"}:
            abort(400, "记录状态不正确。")
        for key in ("from", "to"):
            if filters[key]:
                try:
                    date.fromisoformat(filters[key])
                except ValueError:
                    abort(400, "日期格式不正确。")
        if filters["from"] and filters["to"] and filters["from"] > filters["to"]:
            abort(400, "开始日期不能晚于结束日期。")
        return filters

    @app.get("/ledger")
    @login_required
    def ledger():
        from .ledger import KIND_LABELS, ledger_currency_summary, ledger_query, ledger_totals
        filters = ledger_filters_from_request()
        sql, params = ledger_query(filters)
        total = connection().execute(f"SELECT COUNT(*) FROM ({sql})", params).fetchone()[0]
        per_page = 50
        pages = max(1, math.ceil(total / per_page))
        page = min(max(1, request.args.get("page", 1, type=int) or 1), pages)
        rows = connection().execute(sql + " LIMIT ? OFFSET ?", params + [per_page, (page - 1) * per_page]).fetchall()
        summary = ledger_currency_summary(ledger_totals(connection(), filters))
        active = {key: value for key, value in filters.items() if value}
        return render_template("ledger.html", rows=rows, summary=summary, filters=filters, total=total, page=page, pages=pages,
                               query=urlencode(active), previous=urlencode({**active, "page": page - 1}),
                               following=urlencode({**active, "page": page + 1}), kind_labels=KIND_LABELS, **db.choices(connection()))

    @app.get("/ledger/export.xlsx")
    @login_required
    def export_ledger():
        from .ledger import ledger_rows, ledger_totals, ledger_workbook
        filters = ledger_filters_from_request()
        rows = ledger_rows(connection(), filters)
        totals = ledger_totals(connection(), filters)
        contract_filters = {key: filters[key] for key in ("q", "company", "currency") if filters[key]}
        contract_sql, contract_params = db.document_query(contract_filters)
        contract_rows = connection().execute(contract_sql, contract_params).fetchall()
        workbook = ledger_workbook(rows, contract_rows, totals, generated_at=datetime.now().astimezone().strftime("%Y-%m-%d %H:%M"))
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        with db.transaction(root) as conn:
            db.audit(conn, "ledger_export", user_id=g.user["id"], details={"count": len(rows), "filters": filters})
        return send_file(output, as_attachment=True, download_name="收付款台账.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def file_filters_from_request():
        filters = {key: request.args.get(key, "").strip()[:200] for key in ("q", "company", "year", "status", "filetype")}
        if filters["company"] and (not filters["company"].isdigit() or len(filters["company"]) > 18):
            abort(400, "公司筛选不正确。")
        if filters["year"] and not re.fullmatch(r"\d{4}", filters["year"]):
            abort(400, "年份筛选不正确。")
        if filters["status"] and filters["status"] not in {"active", "archived"}:
            abort(400, "归档状态不正确。")
        if filters["filetype"] and filters["filetype"] not in {"pdf", "word", "image", "other"}:
            abort(400, "文件类型不正确。")
        return filters

    @app.get("/files")
    @login_required
    def files():
        from .file_archive import file_query, file_summary
        filters = file_filters_from_request()
        sql, params = file_query(filters)
        rows = connection().execute(sql, params).fetchall()
        summary = file_summary(rows)
        years = [row[0] for row in connection().execute(
            "SELECT DISTINCT substr(COALESCE(signed_date, created_at), 1, 4) AS value FROM documents ORDER BY value DESC").fetchall() if row[0]]
        return render_template("files.html", rows=rows, filters=filters, summary=summary, years=years,
                               query=urlencode({key: value for key, value in filters.items() if value}), **db.choices(connection()))

    @app.post("/files/bulk")
    @login_required
    def files_bulk():
        from .file_archive import bulk_set_archived
        action = request.form.get("action")
        if action not in {"archive", "restore"}:
            abort(400, "批量操作不正确。")
        identifiers = [item for item in request.form.getlist("ids") if item][:500]
        if not identifiers:
            abort(400, "请先选择文件。")
        with db.transaction(root) as conn:
            result = bulk_set_archived(conn, identifiers, archived=action == "archive", user_id=g.user["id"], created_at=db.now())
        message = f"已{'归档' if action == 'archive' else '恢复'} {len(result['updated'])} 份文件。"
        if result["skipped"]:
            message += f" {len(result['skipped'])} 份状态未变化，已跳过。"
        if result["conflicts"]:
            message += f" {len(result['conflicts'])} 份与在库原件重复，未能恢复。"
        if result["missing"]:
            message += f" {len(result['missing'])} 份文件记录不存在。"
        flash(message, "success" if result["updated"] else "info")
        query = {key: value for key, value in request.form.items() if key in {"q", "company", "year", "status", "filetype"} and value}
        return redirect(url_for("files", **query))

    @app.get("/files/export.zip")
    @login_required
    def export_files():
        from .file_archive import build_archive_zip, file_query
        filters = file_filters_from_request()
        sql, params = file_query(filters)
        rows = connection().execute(sql, params).fetchall()
        labels = {"q": "关键词", "company": "公司", "year": "年份", "status": "归档状态", "filetype": "文件类型"}
        filter_text = "；".join(f"{labels[key]}：{value}" for key, value in filters.items() if value) or "全部文件"
        try:
            output = build_archive_zip(root, rows, generated_at=datetime.now().astimezone().strftime("%Y-%m-%d %H:%M"), filter_text=filter_text)
        except ValueError as exc:
            abort(409, str(exc))
        with db.transaction(root) as conn:
            db.audit(conn, "file_archive_export", user_id=g.user["id"], details={"count": len(rows), "filters": filters})
        response = send_file(output, as_attachment=True, download_name=f"文件归档-{date.today().isoformat()}.zip", mimetype="application/zip")
        response.call_on_close(output.close)
        return response

    @app.get("/export.zip")
    @login_required
    def export_zip():
        rows = filtered_rows()
        output = tempfile.SpooledTemporaryFile(max_size=20 * 1024 * 1024, mode="w+b")
        try:
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("合同清单.xlsx", excel_content(rows).getvalue())
                for row in rows:
                    path = root / "originals" / row["stored_filename"]
                    if not path.is_file():
                        abort(409, "导出中有原件缺失，请联系管理员恢复后重试。")
                    archive.write(path, f"原件/{row['id'][:8]}_{row['original_filename']}")
            output.seek(0)
            with db.transaction(root) as conn:
                db.audit(conn, "export_zip", user_id=g.user["id"], details={"count": len(rows), "filters": dict(request.args)})
            response = send_file(output, as_attachment=True, download_name="合同原件.zip", mimetype="application/zip")
            response.call_on_close(output.close)
            return response
        except Exception:
            output.close()
            raise

    @app.get("/admin")
    @admin_required
    def admin():
        users = connection().execute("SELECT id, username, role, active, created_at FROM users ORDER BY id").fetchall()
        events = connection().execute("SELECT a.*, u.username FROM audit_events a LEFT JOIN users u ON u.id = a.user_id ORDER BY a.id DESC LIMIT 100").fetchall()
        return render_template("admin.html", users=users, events=events, **db.choices(connection()))

    @app.post("/admin/restart")
    @admin_required
    def restart_service():
        if not app.config.get('SERVICE_PORT'):
            abort(503, "当前启动方式不支持网页重启，请使用项目中的重启脚本。")
        script = PROJECT_ROOT / "restart.ps1"
        if not script.is_file():
            abort(503, "找不到重启脚本。")
        subprocess.Popen(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), '-Port', str(app.config['SERVICE_PORT']), '-ListenAddress', app.config['SERVICE_HOST']], cwd=str(PROJECT_ROOT), creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return "服务正在重启，请稍候刷新页面。", 202

    @app.post("/admin/users")
    @admin_required
    def add_user():
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        role = request.form.get("role", "user")
        if not re.fullmatch(r"[\w\-.]{2,40}", username) or len(password) < 10 or role not in {"admin", "user"}:
            abort(400, "账号为 2–40 位文字、数字或下划线，密码至少 10 位。")
        try:
            with db.transaction(root) as conn:
                conn.execute("INSERT INTO users(username, password_hash, role, created_at) VALUES (?, ?, ?, ?)", (username, password_hasher.hash(password), role, db.now()))
                db.audit(conn, "user_create", user_id=g.user["id"], details={"username": username, "role": role})
        except sqlite3.IntegrityError:
            flash("账号名称已存在。", "error")
        else:
            flash("账号已创建。", "success")
        return redirect(url_for("admin"))

    @app.post("/admin/users/<int:user_id>")
    @admin_required
    def update_user(user_id):
        role = request.form.get("role", "user")
        active = int(request.form.get("active") == "1")
        password = request.form.get("password", "")
        if role not in {"admin", "user"} or (password and len(password) < 10):
            abort(400, "角色不正确，或新密码不足 10 位。")
        with db.transaction(root) as conn:
            user = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
            if not user:
                abort(404)
            if user["role"] == "admin" and user["active"] and (not active or role != "admin"):
                admins = conn.execute("SELECT COUNT(*) FROM users WHERE role = 'admin' AND active = 1").fetchone()[0]
                if admins <= 1:
                    abort(400, "必须保留至少一个启用的管理员。")
            conn.execute("UPDATE users SET role = ?, active = ?, password_hash = ?, session_version = session_version + 1 WHERE id = ?",
                         (role, active, password_hasher.hash(password) if password else user["password_hash"], user_id))
            db.audit(conn, "user_update", user_id=g.user["id"], details={"target_user": user["username"], "role": role, "active": active, "password_reset": bool(password)})
        flash("账号已更新，该账号需重新登录。", "success")
        return redirect(url_for("admin"))

    @app.post("/admin/taxonomy")
    @admin_required
    def taxonomy():
        kind = request.form.get("kind")
        name = request.form.get("name", "").strip()
        identifier = request.form.get("id", "")
        type_id = request.form.get("type_id", "")
        if kind not in {"type", "subtype"} or not name or len(name) > 60:
            abort(400, "分类名称需为 1–60 个字符。")
        try:
            with db.transaction(root) as conn:
                if kind == "type":
                    if identifier:
                        conn.execute("UPDATE import_types SET name = ? WHERE id = ?", (name, identifier))
                    else:
                        conn.execute("INSERT INTO import_types(name) VALUES (?)", (name,))
                else:
                    if identifier:
                        conn.execute("UPDATE subtypes SET name = ? WHERE id = ?", (name, identifier))
                    else:
                        if not conn.execute("SELECT 1 FROM import_types WHERE id = ?", (type_id,)).fetchone():
                            abort(400, "请先选择导入类型。")
                        conn.execute("INSERT INTO subtypes(type_id, name) VALUES (?, ?)", (type_id, name))
                db.audit(conn, "taxonomy", user_id=g.user["id"], details={"kind": kind, "name": name})
        except sqlite3.IntegrityError:
            flash("同一层级已存在这个名称。", "error")
        else:
            flash("分类已保存。", "success")
        return redirect(url_for("admin"))

    @app.post("/admin/backup")
    @admin_required
    def backup():
        try:
            output = create_backup(root)
        except (ValueError, OSError) as exc:
            abort(409, str(exc))
        with db.transaction(root) as conn:
            db.audit(conn, "backup", user_id=g.user["id"], details={"filename": output.name})
        return send_file(output, as_attachment=True, download_name=output.name)

    @app.get("/health")
    def health():
        connection().execute("SELECT 1").fetchone()
        return jsonify(status="ok")

    @app.errorhandler(400)
    @app.errorhandler(403)
    @app.errorhandler(404)
    @app.errorhandler(409)
    @app.errorhandler(413)
    @app.errorhandler(429)
    def error_page(error):
        message = "文件总大小超过 100 MB，请分批上传。" if error.code == 413 else error.description
        return render_template("error.html", code=error.code, message=message), error.code

    return app
