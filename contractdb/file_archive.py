from __future__ import annotations

import io
import json
import re
import sqlite3
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="183B39")
HEADER_FONT = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
BODY_FONT = Font(name="Microsoft YaHei", size=10)
MONEY_FORMAT = "#,##0.00;[Red](#,##0.00);0.00"
ARCHIVE_MANIFEST = "文件归档清单.xlsx"
ARCHIVE_NOTE = "归档说明.txt"

_SUFFIXES = {
    "pdf": (".pdf",),
    "word": (".doc", ".docx"),
    "image": (".png", ".jpg", ".jpeg", ".tif", ".tiff"),
}


def _value(row, key, default=None):
    try:
        return row[key]
    except (KeyError, IndexError, TypeError):
        return default


def _escape_like(value: str) -> str:
    return value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def _suffix_conditions(column: str, suffixes: tuple[str, ...]) -> str:
    return " OR ".join(f"lower({column}) LIKE '%{suffix}'" for suffix in suffixes)


def _safe_segment(value: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", value or "").strip(" .")
    return cleaned or "未命名"


def file_query(filters: dict) -> tuple[str, list]:
    """Return the SQL and parameters for the file archive listing (no LIMIT)."""
    sql = """
      SELECT d.*, c.name AS company_name, p.name AS project_name,
             t.name AS type_name, s.name AS subtype_name,
             COALESCE(f.received_minor, 0) AS received_minor,
             COALESCE(f.invoiced_minor, 0) AS invoiced_minor,
             COALESCE(f.paid_minor, 0) AS paid_minor
      FROM documents d
      LEFT JOIN companies c ON c.id = d.company_id
      LEFT JOIN projects p ON p.id = d.project_id
      LEFT JOIN import_types t ON t.id = d.type_id
      LEFT JOIN subtypes s ON s.id = d.subtype_id
      LEFT JOIN (
        SELECT document_id,
          SUM(CASE WHEN kind = 'receipt' THEN amount_minor ELSE 0 END) AS received_minor,
          SUM(CASE WHEN kind = 'invoice' THEN amount_minor ELSE 0 END) AS invoiced_minor,
          SUM(CASE WHEN kind = 'payment' THEN amount_minor ELSE 0 END) AS paid_minor
        FROM finance_records WHERE voided_at IS NULL GROUP BY document_id
      ) f ON f.document_id = d.id
    """
    clauses: list[str] = []
    params: list = []
    keyword = (filters.get("q") or "").strip()
    if keyword:
        pattern = f"%{_escape_like(keyword)}%"
        clauses.append(
            "(d.original_filename LIKE ? ESCAPE '\\' OR d.title LIKE ? ESCAPE '\\' "
            "OR d.contract_number LIKE ? ESCAPE '\\' OR c.name LIKE ? ESCAPE '\\' "
            "OR p.name LIKE ? ESCAPE '\\')"
        )
        params.extend([pattern] * 5)
    if filters.get("company"):
        clauses.append("d.company_id = ?")
        params.append(int(filters["company"]))
    if filters.get("year"):
        clauses.append("substr(COALESCE(d.signed_date, d.created_at), 1, 4) = ?")
        params.append(str(filters["year"]))
    if filters.get("status") == "active":
        clauses.append("d.archived_at IS NULL")
    elif filters.get("status") == "archived":
        clauses.append("d.archived_at IS NOT NULL")
    filetype = filters.get("filetype")
    if filetype in _SUFFIXES:
        clauses.append(f"({_suffix_conditions('d.original_filename', _SUFFIXES[filetype])})")
    elif filetype == "other":
        known = " OR ".join(_suffix_conditions("d.original_filename", suffixes) for suffixes in _SUFFIXES.values())
        clauses.append(f"NOT ({known})")
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    sql += " ORDER BY (d.archived_at IS NOT NULL), d.created_at DESC, d.id DESC"
    return sql, params


def file_type_label(row) -> str:
    lowered = str(_value(row, "original_filename", "") or "").lower()
    if lowered.endswith(".pdf"):
        return "PDF"
    if lowered.endswith((".doc", ".docx")):
        return "Word"
    if lowered.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff")):
        return "图片"
    return "其他"


def file_summary(rows) -> dict:
    total_size = 0
    archived = 0
    for row in rows:
        total_size += row["file_size"] or 0
        if row["archived_at"]:
            archived += 1
    return {
        "total_count": len(rows),
        "active_count": len(rows) - archived,
        "archived_count": archived,
        "total_size": total_size,
    }


def manifest_workbook(rows, *, generated_at: str, filter_text: str) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文件归档清单"
    sheet.append([
        "序号", "文件状态", "原文件名", "合同名称", "公司", "合同编号", "定点项目",
        "文件类型", "文件大小(字节)", "签订日期", "导入时间", "归档时间", "SHA256", "存储文件名",
    ])
    for index, row in enumerate(rows, 1):
        sheet.append([
            index,
            "已归档" if _value(row, "archived_at") else "在库",
            _value(row, "original_filename", "") or "",
            _value(row, "title") or "",
            _value(row, "company_name") or "",
            _value(row, "contract_number") or "",
            _value(row, "project_name") or "",
            file_type_label(row),
            _value(row, "file_size") or 0,
            _value(row, "signed_date") or "",
            _value(row, "created_at") or "",
            _value(row, "archived_at") or "",
            _value(row, "sha256") or "",
            _value(row, "stored_filename") or "",
        ])
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            if not isinstance(cell.value, (int, float)):
                cell.data_type = "s"
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = (6, 10, 50, 40, 34, 25, 30, 10, 14, 13, 20, 20, 68, 40)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:N{len(rows) + 1}"
    sheet.append([])
    sheet.append([f"生成时间：{generated_at}", f"筛选条件：{filter_text}", f"文件总数：{len(rows)}"])
    for cell in sheet[sheet.max_row]:
        cell.font = BODY_FONT
    workbook.properties.title = "文件归档清单"
    workbook.properties.description = f"{filter_text} · 生成时间 {generated_at}"
    return workbook


def _workbook_bytes(workbook: Workbook) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_archive_zip(root: Path, rows, *, generated_at: str, filter_text: str) -> tempfile.SpooledTemporaryFile:
    """Package the manifest plus every original file currently matching the filter."""
    base = Path(root)
    entries: list[tuple[str, object]] = []
    used: set[str] = set()
    lines: list[str] = []
    total_size = 0
    for index, row in enumerate(rows, 1):
        company = _safe_segment(row["company_name"] or "未确认公司")
        year = ((row["signed_date"] or row["created_at"] or "")[:4]) or "日期待补齐"
        filename = row["original_filename"]
        entry = f"文件/{company}/{year}/{index:02d}_{filename}"
        if entry in used:
            entry = f"文件/{company}/{year}/{index:02d}_{row['id'][:8]}_{filename}"
        used.add(entry)
        path = base / "originals" / row["stored_filename"]
        if not path.is_file():
            raise ValueError(f"原件缺失：{row['original_filename']}")
        size = path.stat().st_size
        total_size += size
        entries.append((entry, path))
        lines.append(f"{index}. {filename} -> {entry}（{size} 字节）")
    note = "\n".join([
        "文件归档说明",
        f"生成时间：{generated_at}",
        f"筛选条件：{filter_text}",
        f"文件总数：{len(rows)}",
        f"文件总大小：{total_size} 字节",
        "",
        "SHA256 列可核对文件完整性，存储文件名为数据库内部名称。",
        "",
        "文件清单：",
        *lines,
        "",
    ])
    output = tempfile.SpooledTemporaryFile(max_size=20 * 1024 * 1024, mode="w+b")
    try:
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(ARCHIVE_MANIFEST, _workbook_bytes(manifest_workbook(rows, generated_at=generated_at, filter_text=filter_text)))
            archive.writestr(ARCHIVE_NOTE, note.encode("utf-8"))
            for entry, path in entries:
                archive.write(path, entry)
    except Exception:
        output.close()
        raise
    output.seek(0)
    return output


def bulk_set_archived(connection, identifiers: list[str], *, archived: bool, user_id: int, created_at: str) -> dict:
    """Archive or restore many documents at once, keeping per-item failures isolated."""
    result: dict[str, list[str]] = {"updated": [], "skipped": [], "conflicts": [], "missing": []}
    seen: set[str] = set()
    event = "archive" if archived else "restore"
    for identifier in identifiers:
        if not identifier or identifier in seen:
            continue
        seen.add(identifier)
        connection.execute("SAVEPOINT file_archive_item")
        try:
            row = connection.execute("SELECT archived_at FROM documents WHERE id = ?", (identifier,)).fetchone()
            if not row:
                result["missing"].append(identifier)
                connection.execute("RELEASE file_archive_item")
                continue
            if (row[0] is not None) == archived:
                result["skipped"].append(identifier)
                connection.execute("RELEASE file_archive_item")
                continue
            connection.execute(
                "UPDATE documents SET archived_at = ?, updated_at = ?, revision = revision + 1 WHERE id = ?",
                (created_at if archived else None, created_at, identifier),
            )
            connection.execute(
                "INSERT INTO audit_events(document_id, user_id, event, details_json, created_at) VALUES (?, ?, ?, ?, ?)",
                (identifier, user_id, event, json.dumps({"source": "file_archive", "batch": True}, ensure_ascii=False), created_at),
            )
            result["updated"].append(identifier)
            connection.execute("RELEASE file_archive_item")
        except sqlite3.IntegrityError:
            connection.execute("ROLLBACK TO file_archive_item")
            connection.execute("RELEASE file_archive_item")
            result["conflicts"].append(identifier)
        except Exception:
            connection.execute("ROLLBACK TO file_archive_item")
            connection.execute("RELEASE file_archive_item")
            raise
    if result["updated"]:
        connection.execute(
            "INSERT INTO audit_events(document_id, user_id, event, details_json, created_at) VALUES (NULL, ?, ?, ?, ?)",
            (user_id, "file_archive_batch" if archived else "file_restore_batch",
             json.dumps({"count": len(result["updated"])}, ensure_ascii=False), created_at),
        )
    return result
