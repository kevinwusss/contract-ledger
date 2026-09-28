from __future__ import annotations

import re
from datetime import date
from decimal import Decimal
from pathlib import Path
from functools import lru_cache

from docx import Document
from PIL import Image
from pypdf import PdfReader


MAX_TEXT_PAGES = 30
MAX_OCR_PAGES = 6
MAX_INDEX_CHARS = 200_000


@lru_cache(maxsize=1)
def _ocr_engine():
    from rapidocr_onnxruntime import RapidOCR
    return RapidOCR()


def _ocr_image(image) -> str:
    import numpy as np
    engine = _ocr_engine()
    result, _ = engine(np.asarray(image.convert("RGB")))
    if not result:
        return ""
    return "\n".join(str(item[1]) for item in result if len(item) > 1)


def extract_pages(path: Path) -> tuple[list[tuple[str, str]], str]:
    suffix = path.suffix.lower()
    pages: list[tuple[str, str]] = []
    method = "text"
    if suffix == ".pdf":
        reader = PdfReader(str(path), strict=False)
        count = min(len(reader.pages), MAX_TEXT_PAGES)
        sparse: list[int] = []
        for index in range(count):
            text = reader.pages[index].extract_text() or ""
            pages.append((f"第 {index + 1} 页", text))
            if len(sparse) < MAX_OCR_PAGES and len(text.strip()) < 80:
                sparse.append(index)
        if sparse:
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(path))
            try:
                for index in reversed(sparse):
                    bitmap = pdf[index].render(scale=1.7)
                    image = bitmap.to_pil()
                    ocr_text = _ocr_image(image)
                    if ocr_text:
                        pages.insert(index + 1, (f"第 {index + 1} 页（OCR）", ocr_text))
                        method = "ocr+text"
            finally:
                pdf.close()
    elif suffix == ".docx":
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                lines.append("  ".join(cell.text for cell in row.cells))
        pages.append(("Word 正文", "\n".join(lines)))
    else:
        with Image.open(path) as image:
            pages.append(("图片 OCR", _ocr_image(image)))
        method = "ocr"
    return pages, method


COMPANY_RE = re.compile(
    r"(?:甲方|乙方|委托方|受托方|采购方|供应商|客户|合同相对方)\s*[：:，,]?\s*"
    r"([\u4e00-\u9fa5A-Za-z0-9（）()·\-]{3,55}?(?:股份有限公司|有限责任公司|有限公司|集团公司))"
)
NUMBER_RE = re.compile(r"(?:合同编号|合同号|协议编号|协议号)\s*[：:]?\s*([A-Za-z0-9][A-Za-z0-9\-/_.]{3,60})")
PROJECT_RE = re.compile(r"(?:项目名称|定点项目)\s*[：:]\s*([^\n\r，,；;]{2,80})")
DATE_RE = re.compile(r"(?:签订日期|签署日期|签约日期)\s*[：:]?\s*(20\d{2})[年./-](\d{1,2})[月./-](\d{1,2})日?")
AMOUNT_RE = re.compile(
    r"(?:合同总金额|合同总价|合同金额|含税总金额|含税总价|项目总金额)\s*"
    r"(?:[（(]含税[)）])?\s*[：:为是]?\s*(?:人民币|RMB|CNY|[¥￥])?\s*"
    r"([0-9][0-9,，]*(?:\.[0-9]{1,2})?)\s*(万元|万|元)?"
)


def suggest_fields(pages: list[tuple[str, str]], method: str) -> dict[str, list[dict[str, str]]]:
    suggestions: dict[str, list[dict[str, str]]] = {
        "company": [], "contract_number": [], "project": [], "signed_date": [],
        "subtype": [], "amount_hint": []
    }
    patterns = {
        "company": COMPANY_RE,
        "contract_number": NUMBER_RE,
        "project": PROJECT_RE,
        "signed_date": DATE_RE,
    }
    for label, text in pages:
        for field, pattern in patterns.items():
            for match in list(pattern.finditer(text))[:5]:
                if field == "signed_date":
                    year, month, day = (int(value) for value in match.groups())
                    try:
                        date(year, month, day)
                    except ValueError:
                        continue
                    value = f"{year:04d}-{month:02d}-{day:02d}"
                else:
                    value = match.group(1).strip(" ：:，,。 ")
                if not value or any(item["value"] == value for item in suggestions[field]):
                    continue
                snippet = re.sub(r"\s+", " ", text[max(0, match.start() - 20):match.end() + 30]).strip()
                source_method = "ocr" if "OCR" in label else "text"
                suggestions[field].append({"value": value, "page": label, "excerpt": snippet, "method": source_method})
                if len(suggestions[field]) >= 5:
                    break
        for match in list(re.finditer("开发费", text))[:1]:
            if not suggestions["subtype"]:
                snippet = re.sub(r"\s+", " ", text[max(0, match.start() - 20):match.end() + 30]).strip()
                suggestions["subtype"].append({"value": "开发费", "page": label, "excerpt": snippet,
                                               "method": "ocr" if "OCR" in label else "text"})
        for match in list(AMOUNT_RE.finditer(text))[:5]:
            number = match.group(1).replace(",", "").replace("，", "")
            amount = Decimal(number) * (10000 if match.group(2) in {"万", "万元"} else 1)
            value = f"{amount:.2f}"
            if any(item["value"] == value for item in suggestions["amount_hint"]):
                continue
            snippet = re.sub(r"\s+", " ", text[max(0, match.start() - 20):match.end() + 30]).strip()
            suggestions["amount_hint"].append({"value": value, "page": label, "excerpt": snippet,
                                               "method": "ocr" if "OCR" in label else "text"})
            if len(suggestions["amount_hint"]) >= 5:
                break
    return suggestions


def extract_document(path: Path) -> tuple[str, dict[str, list[dict[str, str]]]]:
    pages, method = extract_pages(path)
    text = "\n".join(f"[{label}]\n{content}" for label, content in pages)[:MAX_INDEX_CHARS]
    return text, suggest_fields(pages, method)
