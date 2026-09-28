from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .finance import CONFIRMATION_LABELS, confirmation_status


KIND_LABELS = {"receipt": "我方收款", "payment": "我方付款", "invoice": "我方开票"}
MONEY_FORMAT = "#,##0.00;[Red](#,##0.00);0.00"
HEADER_FILL = PatternFill("solid", fgColor="183B39")
HEADER_FONT = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
BODY_FONT = Font(name="Microsoft YaHei", size=10)

LEDGER_COLUMNS = ("日期", "类型", "公司", "合同名称", "合同编号", "定点项目", "币种", "金额",
                  "凭证 / 发票号", "具体说明", "状态", "登记人", "登记时间")
SUMMARY_COLUMNS = ("币种", "收款合计", "付款合计", "开票合计", "有效笔数", "作废笔数")
CONTRACT_COLUMNS = ("公司", "合同名称", "合同编号", "定点项目", "币种", "对方应付我方", "我方已收", "我方未收",
                    "我方应付对方", "我方已付", "我方未付", "我方已开票", "我方未开票", "款项确认")


def _value(row, key, default=None):
    try:
        return row[key]
    except (KeyError, IndexError, TypeError):
        return default


def _money(value):
    return None if value is None else value / 100


def _escape_like(value: str) -> str:
    return value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


FROM_CLAUSE = """
      FROM finance_records f
      JOIN documents d ON d.id = f.document_id
      LEFT JOIN companies c ON c.id = d.company_id
      LEFT JOIN projects p ON p.id = d.project_id
      LEFT JOIN users u ON u.id = f.created_by
    """
SELECT_CLAUSE = """
      SELECT f.id, f.document_id, f.kind, f.amount_minor, f.occurred_on, f.reference, f.description,
             f.created_at, f.voided_at, f.void_reason,
             d.original_filename, d.title, d.contract_number, d.currency, d.signed_date, d.archived_at,
             c.name AS company_name, p.name AS project_name, u.username AS creator_name
    """


def _conditions(filters: dict) -> tuple[list[str], list]:
    """Shared WHERE clauses so the listing and the totals can never drift apart."""
    clauses: list[str] = []
    params: list = []
    keyword = (filters.get("q") or "").strip()
    if keyword:
        pattern = f"%{_escape_like(keyword)}%"
        clauses.append(
            "(d.title LIKE ? ESCAPE '\\' OR d.original_filename LIKE ? ESCAPE '\\' "
            "OR d.contract_number LIKE ? ESCAPE '\\' OR c.name LIKE ? ESCAPE '\\' "
            "OR p.name LIKE ? ESCAPE '\\' OR f.reference LIKE ? ESCAPE '\\' "
            "OR f.description LIKE ? ESCAPE '\\')"
        )
        params.extend([pattern] * 7)
    if filters.get("company"):
        clauses.append("d.company_id = ?")
        params.append(int(filters["company"]))
    if filters.get("kind"):
        clauses.append("f.kind = ?")
        params.append(filters["kind"])
    if filters.get("currency"):
        clauses.append("d.currency = ?")
        params.append(filters["currency"])
    if filters.get("from"):
        clauses.append("f.occurred_on >= ?")
        params.append(filters["from"])
    if filters.get("to"):
        clauses.append("f.occurred_on <= ?")
        params.append(filters["to"])
    if filters.get("status") == "valid":
        clauses.append("f.voided_at IS NULL")
    elif filters.get("status") == "void":
        clauses.append("f.voided_at IS NOT NULL")
    if filters.get("document"):
        clauses.append("f.document_id = ?")
        params.append(filters["document"])
    return clauses, params


def ledger_query(filters: dict) -> tuple[str, list]:
    """Return the SQL and parameters for every matching finance record (no LIMIT)."""
    clauses, params = _conditions(filters)
    sql = SELECT_CLAUSE + FROM_CLAUSE
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    sql += " ORDER BY f.occurred_on DESC, f.id DESC"
    return sql, params


def _totals_query(filters: dict) -> tuple[str, list]:
    clauses, params = _conditions(filters)
    clauses.append("f.voided_at IS NULL")
    sql = ("SELECT d.currency AS currency, f.kind AS kind, SUM(f.amount_minor) AS total_minor, COUNT(*) AS count"
           + FROM_CLAUSE + " WHERE " + " AND ".join(clauses)
           + " GROUP BY d.currency, f.kind ORDER BY d.currency, f.kind")
    return sql, params


def ledger_rows(connection, filters: dict) -> list:
    sql, params = ledger_query(filters)
    return connection.execute(sql, params).fetchall()


def ledger_totals(connection, filters: dict) -> list[dict]:
    sql, params = _totals_query(filters)
    return [{"currency": row["currency"], "kind": row["kind"],
             "total_minor": row["total_minor"] or 0, "count": row["count"]} for row in connection.execute(sql, params).fetchall()]


def ledger_currency_summary(totals: list[dict]) -> list[dict]:
    groups: dict[str, dict] = {}
    for item in totals:
        group = groups.setdefault(item["currency"], {
            "currency": item["currency"], "receipt_minor": 0, "payment_minor": 0, "invoice_minor": 0,
            "receipt_count": 0, "payment_count": 0, "invoice_count": 0, "total_count": 0,
        })
        group[f"{item['kind']}_minor"] += item["total_minor"]
        group[f"{item['kind']}_count"] += item["count"]
        group["total_count"] += item["count"]
    return [groups[key] for key in sorted(groups)]


def _apply_body_style(sheet) -> None:
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            if isinstance(cell.value, (int, float)):
                cell.number_format = MONEY_FORMAT
            else:
                cell.data_type = "s"
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _apply_header(sheet, widths) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def ledger_workbook(rows, contract_rows, totals, *, generated_at: str) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "收付款台账"
    sheet.append(list(LEDGER_COLUMNS))
    void_counts: dict[str, int] = {}
    for row in rows:
        if row["voided_at"]:
            void_counts[row["currency"]] = void_counts.get(row["currency"], 0) + 1
        description = row["description"] or ""
        if row["voided_at"] and row["void_reason"]:
            description = f"{description} 作废原因：{row['void_reason']}".strip()
        sheet.append([
            row["occurred_on"],
            KIND_LABELS.get(row["kind"], row["kind"]),
            row["company_name"] or "",
            row["title"] or row["original_filename"] or "",
            row["contract_number"] or "",
            row["project_name"] or "",
            row["currency"],
            _money(row["amount_minor"]),
            row["reference"] or "",
            description,
            "已作废" if row["voided_at"] else "有效",
            row["creator_name"] or "系统",
            row["created_at"] or "",
        ])
    _apply_header(sheet, (13, 12, 34, 40, 25, 30, 10, 16, 22, 48, 10, 14, 22))
    _apply_body_style(sheet)

    summary = workbook.create_sheet("按币种汇总")
    summary.append(list(SUMMARY_COLUMNS))
    for group in ledger_currency_summary(totals):
        summary.append([group["currency"], _money(group["receipt_minor"]), _money(group["payment_minor"]),
                        _money(group["invoice_minor"]), group["total_count"], void_counts.get(group["currency"], 0)])
    for currency, count in sorted(void_counts.items()):
        if currency not in {group["currency"] for group in ledger_currency_summary(totals)}:
            summary.append([currency, 0, 0, 0, 0, count])
    _apply_header(summary, (12, 18, 18, 18, 12, 12))
    _apply_body_style(summary)

    contracts = workbook.create_sheet("合同台账")
    contracts.append(list(CONTRACT_COLUMNS))
    for row in contract_rows:
        amount = row["amount_minor"]
        payable = row["payable_minor"]
        received = row["received_minor"] or 0
        paid = row["paid_minor"] or 0
        invoiced = row["invoiced_minor"] or 0
        contracts.append([
            row["company_name"] or "",
            row["title"] or row["original_filename"] or "",
            row["contract_number"] or "",
            row["project_name"] or "",
            row["currency"],
            _money(amount),
            _money(received),
            None if amount is None else _money(max(amount - received, 0)),
            _money(payable),
            _money(paid),
            None if payable is None else _money(max(payable - paid, 0)),
            _money(invoiced),
            None if amount is None else _money(max(amount - invoiced, 0)),
            CONFIRMATION_LABELS[confirmation_status(row)],
        ])
    _apply_header(contracts, (34, 40, 25, 30, 10, 18, 16, 16, 18, 16, 16, 16, 16, 16))
    _apply_body_style(contracts)

    workbook.properties.title = "收付款台账"
    workbook.properties.description = f"生成时间 {generated_at}"
    return workbook
